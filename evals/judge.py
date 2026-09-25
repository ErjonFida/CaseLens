import argparse
import hashlib
import json
import os
import re
import threading
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

import config  # noqa: F401  loads .env, for GEMINI_API_KEY
from . import EVAL_ROOT, REPORT_DIR


JUDGE_MODEL = "gemini-3.7-flash"
GRADES = ("correct", "partial", "wrong", "abstain-correct", "abstain-wrong")
GOOD = {"correct", "abstain-correct"}
CACHE = EVAL_ROOT / ".cache" / "judge.json"

RUBRIC = """You grade answers from a legal assistant against expert annotations of a contract.

Each case gives the question; whether the contract contains the clause asked about; when it does, the annotated clause text, its pages, and whether the assistant's context included an annotated page; and the assistant's answer. You do not see the contract, so judge the answer's claims against the annotation.

Choose exactly one grade.

If the contract CONTAINS the clause:
- correct: the answer states the substance of the annotated clause.
- partial: the substance is incomplete, or partly right and partly wrong, or right but attributed to a page that is not annotated.
- wrong: the answer contradicts the annotated clause, or presents a different provision as the answer.
- abstain-correct: the answer says it cannot find the information, and the context did NOT include an annotated page.
- abstain-wrong: the answer says it cannot find the information although the context included an annotated page.
An answer that says the information is missing and then gives it anyway is graded on what it gives.

If the contract does NOT contain the clause:
- correct: the answer says there is no such provision, or answers no or none, with or without quoting a provision that rules the clause out (a "no third-party beneficiaries" section when asked whether outsiders have rights, for example).
- abstain-correct: the answer only says it cannot find or determine the information.
- wrong: the answer says or implies the clause exists, or offers a different provision as the answer (assignment restrictions offered as what happens on a change of control, for example). Related provisions are fine when the answer first says the asked-for one is absent and labels them as related.
- partial: the answer hedges between saying the clause is absent and asserting it.

Grade the claims, not length, tone or formatting. Reply with JSON only: {"grade": "<grade>", "reason": "<one sentence>"}"""

CORRECTIONS = {
    ("q089", "gemini-40k", "rag10"): "correct",
    ("q089", "gemini-40k", "full"): "correct",
}

_cache_lock = threading.Lock()

_rate_lock = threading.Lock()
_next_start = [0.0]


def _throttle(model: str) -> None:
    per_minute = 20 if "pro" in model else 60
    with _rate_lock:
        now = time.monotonic()
        start = max(now, _next_start[0])
        _next_start[0] = start + 60 / per_minute
    time.sleep(max(0.0, start - now))


def _load_cache() -> dict:
    return json.loads(CACHE.read_text(encoding="utf-8")) if CACHE.exists() else {}


def _save_cache(cache: dict) -> None:
    CACHE.parent.mkdir(parents=True, exist_ok=True)
    CACHE.write_text(json.dumps(cache), encoding="utf-8")


def _case(question: str, clause_absent: bool, spans: list[str], pages: str,
          context_has_page: bool, answer: str) -> str:
    if clause_absent:
        return f"QUESTION: {question}\nCONTRACT CONTAINS THE CLAUSE: no\n\nANSWER:\n{answer}"
    clause = "\n[...]\n".join(s.strip() for s in spans if s.strip())[:3000]
    return (
        f"QUESTION: {question}\nCONTRACT CONTAINS THE CLAUSE: yes\n"
        f"ANNOTATED CLAUSE, PAGE(S) {pages}:\n{clause}\n"
        f"ASSISTANT'S CONTEXT INCLUDED AN ANNOTATED PAGE: {'yes' if context_has_page else 'no'}\n\n"
        f"ANSWER:\n{answer}"
    )


def _ask(model: str, case: str) -> str:
    import google.generativeai as genai

    genai.configure(api_key=os.environ["GEMINI_API_KEY"])
    llm = genai.GenerativeModel(model_name=model, system_instruction=RUBRIC)
    for attempt in range(8):
        _throttle(model)
        try:
            r = llm.generate_content(
                case, generation_config={"temperature": 0, "response_mime_type": "application/json"}
            )
            parts = r.candidates[0].content.parts if r.candidates else []
            return "".join(getattr(p, "text", "") for p in parts if not getattr(p, "thought", False))
        except Exception as e:
            transient = any(code in str(e) for code in ("429", "500", "503", "deadline", "Deadline"))
            if attempt == 7 or not transient:
                raise
            hint = re.search(r"retry in ([0-9.]+)s", str(e))  # the server says how long
            time.sleep(float(hint.group(1)) + 1 if hint else 5 * (attempt + 1))
    return ""


def _parse(raw: str) -> tuple[str, str]:
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        data = {}
    grade = str(data.get("grade", "")).strip().lower()
    if grade not in GRADES:
        found = re.search("|".join(sorted(GRADES, key=len, reverse=True)), raw.lower())
        grade = found.group(0) if found else "unparsed"
    return grade, str(data.get("reason", "")).strip()[:300]


def judge_many(cases: list[dict], model: str = JUDGE_MODEL, workers: int = 8) -> list[dict]:
    """Grade cases (keyword arguments of _case), cached by model and input."""
    cache = _load_cache()
    keys = [hashlib.sha256(json.dumps([model, c], sort_keys=True).encode()).hexdigest() for c in cases]

    def one(i: int) -> None:
        if keys[i] in cache:
            return
        grade, reason = _parse(_ask(model, _case(**cases[i])))
        with _cache_lock:
            cache[keys[i]] = {"model": model, "grade": grade, "reason": reason}

    try:
        with ThreadPoolExecutor(workers) as pool:
            list(pool.map(one, range(len(cases))))
    finally:
        _save_cache(cache)  # keep what finished, even if a call failed
    return [cache[k] for k in keys]


def workbook_cases(path: Path) -> tuple[list[dict], list[dict]]:
    """Rows of the graded workbook, corrected, and the judge cases for them."""
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True)["Answers"]
    header = [c.value for c in ws[1]]
    rows, cases = [], []
    for values in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(header, values))
        if r["category"] == "unanswerable":  # the workbook predates the rename
            r["category"] = "clause_absent"
        r["graded"] = CORRECTIONS.get((r["question_id"], r["report"], r["condition"]), r["Grade"])
        absent = r["category"] == "clause_absent"
        rows.append(r)
        cases.append({
            "question": r["question"],
            "clause_absent": absent,
            "spans": [] if absent else str(r["gold_answer"] or "").split("\n---\n"),
            "pages": "" if absent else str(r["gold_pages"] or ""),
            "context_has_page": r["context_has_page"] == "yes",
            "answer": r["answer"],
        })
    return rows, cases


def agreement(rows: list[dict], verdicts: list[dict]) -> dict:
    graded = [r["graded"] for r in rows]
    judged = [v["grade"] for v in verdicts]

    def good_agreement(idx):
        idx = list(idx)
        return round(sum((graded[i] in GOOD) == (judged[i] in GOOD) for i in idx) / len(idx), 3) if idx else None

    def detection(label):
        g = [label(x) for x in graded]
        j = [label(x) for x in judged]
        tp = sum(a and b for a, b in zip(g, j))
        return {"graded": sum(g), "judged": sum(j), "both": tp}

    by_model = {}
    for m in sorted({r["model"] for r in rows}):
        by_model[m] = good_agreement(i for i, r in enumerate(rows) if r["model"] == m)
    return {
        "cases": len(rows),
        "exact grade agreement": round(sum(a == b for a, b in zip(graded, judged)) / len(rows), 3),
        "good / not good agreement": good_agreement(range(len(rows))),
        "good agreement, clause absent": good_agreement(i for i, r in enumerate(rows) if r["category"] == "clause_absent"),
        "good agreement by answering model": by_model,
        "wrong": detection(lambda x: x == "wrong"),
        "declined": detection(lambda x: x.startswith("abstain")),
        "unparsed": judged.count("unparsed"),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Calibrate the answer judge against graded answers.")
    parser.add_argument("--calibrate", type=Path, required=True, help="The graded workbook")
    parser.add_argument("--model", default=JUDGE_MODEL)
    parser.add_argument("--workers", type=int, default=8)
    args = parser.parse_args()

    rows, cases = workbook_cases(args.calibrate)
    verdicts = judge_many(cases, args.model, args.workers)
    summary = agreement(rows, verdicts)

    print(f"\n  judge {args.model} against {args.calibrate.name} ({summary['cases']} graded answers)\n")
    for key, value in summary.items():
        print(f"  {key:36} {value}")
    confusion = Counter((r["graded"], v["grade"]) for r, v in zip(rows, verdicts))
    print("\n  graded -> judged, where they differ")
    for (g, j), n in sorted(confusion.items(), key=lambda kv: -kv[1]):
        if g != j:
            print(f"    {g:16} -> {j:16} {n}")

    out = REPORT_DIR / f"judge-calibration-{args.model}.json"
    out.write_text(json.dumps({
        "model": args.model,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "workbook": args.calibrate.name,
        "corrections": {"/".join(k): v for k, v in CORRECTIONS.items()},
        "rubric": RUBRIC,
        "summary": summary,
        "rows": [
            {"question_id": r["question_id"], "report": r["report"], "condition": r["condition"],
             "graded": r["graded"], "judged": v["grade"], "reason": v["reason"]}
            for r, v in zip(rows, verdicts)
        ],
    }, indent=2), encoding="utf-8")
    print(f"\n  Written to {out}\n")


if __name__ == "__main__":
    main()
